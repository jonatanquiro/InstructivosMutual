"""Conciliación bancaria: mayor del sistema contra resumen del banco.

Uso:   python conciliar.py COINAG 09/2026

Busca  MAYOR/COINAG092026.pdf  y el resumen del mes en la carpeta del banco,
y deja el resultado en  CONCILIACIONES/COINAG 09-2026.xlsx
"""
import re
import sys
from datetime import date, timedelta
from itertools import combinations, zip_longest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from lectores import BANCOS, ErrorLectura, leer_mayor, leer_resumen, resumenes

BASE = Path(__file__).resolve().parent

DIAS_1A1 = 7           # tolerancia de fechas entre mayor y banco
DIAS_GRUPO = (-5, 3)   # el banco puede acreditar hasta 5 días antes o 3 después del asiento
DIAS_CRUCE_MES = 5     # asientos de los primeros días contra los últimos del mes anterior
TOL_CENTAVOS = 1.00    # diferencia máxima para aceptar un grupo "con diferencia"
TOL_GASTOS = 0.005     # el asiento global de gastos puede diferir hasta 0,5% del resumen

GASTOS = [
    ('Impuesto Ley 25413 (déb/créd)', re.compile(r'LEY\s*(NRO\.?\s*)?25\.?413', re.I)),
    ('SIRCREB (Ingresos Brutos)', re.compile(r'SIRCREB', re.I)),
    ('IVA', re.compile(r'\bI\.?V\.?A\b', re.I)),
    ('Comisiones y mantenimiento', re.compile(r'COMIS|\bCOMI?\b|MANTENIMIENTO', re.I)),
    ('Intereses', re.compile(r'INTERES', re.I)),
    ('Seguros', re.compile(r'SEGUR', re.I)),
]
ASIENTO_GASTOS = re.compile(r'G\.? ?BANC|GASTO')
VALORES = re.compile(r'VALORES AL COBRO')
CUIT_PROPIO = '30619630056'
PROPIAS = re.compile(CUIT_PROPIO + r'|PROPIAS|IGUAL TIT|MISMO TIT|ASOC\w*\.? ?MUT', re.I)
FECHA_EN_CONCEPTO = re.compile(r'^(\d{1,2})/(\d{1,2})\b')


def rubro_gasto(concepto):
    for nombre, rx in GASTOS:
        if rx.search(concepto):
            return nombre
    return None


def parecido(m, b):
    """Cantidad de palabras del concepto del mayor que aparecen en el del banco."""
    texto = (b['concepto'] + ' ' + b['detalle']).upper()
    return sum(1 for p in re.findall(r'[A-ZÑ]{4,}', m['concepto'].upper()) if p in texto)


def subconjunto(objetivo, candidatos, tol=0):
    """Busca entre candidatos [(centavos, item)] un grupo que sume objetivo (± tol)."""
    cands = sorted(candidatos, key=lambda c: -abs(c[0]))[:24]
    resto = [0] * (len(cands) + 1)
    for i in range(len(cands) - 1, -1, -1):
        resto[i] = resto[i + 1] + cands[i][0]
    mejor = None

    def buscar(i, suma, elegidos):
        nonlocal mejor
        if mejor is not None and mejor[0] == 0:
            return
        dif = abs(objetivo - suma)
        if len(elegidos) >= 2 and dif <= tol and (mejor is None or dif < mejor[0]):
            mejor = (dif, list(elegidos))
        if i == len(cands) or len(elegidos) == 10 or suma > objetivo + tol or suma + resto[i] < objetivo - tol:
            return
        elegidos.append(cands[i][1])
        buscar(i + 1, suma + cands[i][0], elegidos)
        elegidos.pop()
        buscar(i + 1, suma, elegidos)

    buscar(0, 0, [])
    return mejor[1] if mejor else None


def conciliar(mayor, banco, banco_ant):
    """Cruza los movimientos. Devuelve la lista de grupos conciliados."""
    for lista in (mayor, banco, banco_ant):
        for x in lista:
            x['ok'] = False
    grupos = []

    def cerrar(tipo, ms, bs, nota=''):
        for x in ms + bs:
            x['ok'] = True
        grupos.append({'tipo': tipo, 'mayor': ms, 'banco': bs, 'nota': nota})

    # 1) Uno a uno: mismo importe, fecha cercana (primero los más cercanos y parecidos)
    pares = []
    for i, m in enumerate(mayor):
        for j, b in enumerate(banco):
            if b['importe'] == m['importe'] and b['importe'] != 0:
                d = abs((b['fecha'] - m['fecha']).days)
                if d <= DIAS_1A1:
                    pares.append((d, -parecido(m, b), i, j, m, b))
    for *_, m, b in sorted(pares, key=lambda p: p[:4]):
        if not m['ok'] and not b['ok']:
            cerrar('Uno a uno', [m], [b])

    # 2) Un asiento de valores contra varias acreditaciones del banco
    for tol in (0, TOL_CENTAVOS):
        for m in mayor:
            if m['ok'] or m['importe'] <= 0:
                continue
            desde = m['fecha'] + timedelta(days=DIAS_GRUPO[0])
            hasta = m['fecha'] + timedelta(days=DIAS_GRUPO[1])
            cands = [(round(b['importe'] * 100), b) for b in banco
                     if not b['ok'] and b['importe'] > 0 and VALORES.search(b['concepto'])
                     and desde <= b['fecha'] <= hasta]
            sel = subconjunto(round(m['importe'] * 100), cands, round(tol * 100))
            if sel:
                dif = round(m['importe'] - sum(b['importe'] for b in sel), 2)
                cerrar('Depósito de valores', [m], sel,
                       f'Diferencia de {dif:,.2f} entre mayor y banco' if dif else '')

    # 2b) Un asiento por varias transferencias entre cuentas propias del mismo día
    for dias in (0, 2):
        for m in mayor:
            if m['ok'] or m['importe'] == 0:
                continue
            signo = 1 if m['importe'] > 0 else -1
            cands = [(round(abs(b['importe']) * 100), b) for b in banco
                     if not b['ok'] and b['importe'] * signo > 0 and abs((b['fecha'] - m['fecha']).days) <= dias
                     and PROPIAS.search(b['concepto'] + ' ' + b['detalle'])]
            sel = subconjunto(round(abs(m['importe']) * 100), cands)
            if sel:
                cerrar('Entre cuentas propias', [m], sel)

    # 2b') Un asiento por varios movimientos iguales del banco (ej. depósitos en efectivo del día)
    for m in mayor:
        if m['ok'] or m['importe'] == 0:
            continue
        signo = 1 if m['importe'] > 0 else -1
        clases = {}
        for b in banco:
            if (not b['ok'] and b['importe'] * signo > 0 and abs((b['fecha'] - m['fecha']).days) <= 1
                    and not rubro_gasto(b['concepto'])):
                clave = re.sub(r'[^A-Za-z ]', '', b['concepto']).split()[:1]
                clases.setdefault(' '.join(clave), []).append((round(abs(b['importe']) * 100), b))
        for cands in clases.values():
            sel = subconjunto(round(abs(m['importe']) * 100), cands)
            if sel:
                cerrar('Varios movimientos', [m], sel)
                break

    # 2c) Un movimiento del banco contra varios asientos (ej. liquidación neta de Cabal:
    #     crédito + débito del día, más pagos y ajustes registrados en días cercanos)
    for antes, despues, tol in ((0, 0, 0), (0, 0, TOL_CENTAVOS), (4, 0, 0), (4, 0, TOL_CENTAVOS),
                                (5, 5, TOL_CENTAVOS)):
        for b in banco:
            if b['ok'] or b['importe'] == 0 or rubro_gasto(b['concepto']):
                continue
            cands = [m for m in mayor if not m['ok']
                     and -antes <= (m['fecha'] - b['fecha']).days <= despues]
            if not 2 <= len(cands) <= 14 or not any(m['fecha'] == b['fecha'] for m in cands):
                continue
            mejor = None
            for n in (2, 3, 4, 5):
                for combo in combinations(cands, n):
                    dif = abs(round(sum(m['importe'] for m in combo) - b['importe'], 2))
                    if dif <= tol + 0.001 and (mejor is None or dif < mejor[0]):
                        mejor = (dif, combo)
                if mejor:
                    break
            if mejor:
                dif = round(sum(m['importe'] for m in mejor[1]) - b['importe'], 2)
                cerrar('Varios asientos', list(mejor[1]), [b],
                       f'Diferencia de {dif:,.2f} entre mayor y banco' if dif else '')

    # 2d) Uno a uno con diferencia de centavos
    for m in mayor:
        if m['ok'] or m['importe'] == 0:
            continue
        cands = [b for b in banco if not b['ok'] and b['importe'] * m['importe'] > 0
                 and abs(b['importe'] - m['importe']) <= TOL_CENTAVOS and abs((b['fecha'] - m['fecha']).days) <= 3]
        if len(cands) == 1:
            dif = round(m['importe'] - cands[0]['importe'], 2)
            cerrar('Uno a uno', [m], cands, f'Diferencia de {dif:,.2f} entre mayor y banco')

    # 2e) Mismo importe con fechas distantes, si no hay otro candidato posible
    for m in mayor:
        if m['ok'] or m['importe'] == 0:
            continue
        cands = [b for b in banco if not b['ok'] and b['importe'] == m['importe']]
        otros = [x for x in mayor if not x['ok'] and x['importe'] == m['importe']]
        if len(cands) == 1 and len(otros) == 1:
            dias = abs((cands[0]['fecha'] - m['fecha']).days)
            cerrar('Uno a uno', [m], cands, f'Fechas distantes: {dias} días entre mayor y banco')

    # 3) Asientos de los primeros días contra el resumen del mes anterior
    if banco_ant:
        inicio = date(mayor[0]['fecha'].year, mayor[0]['fecha'].month, 1) if mayor else None
        # asientos que dicen en el concepto la fecha del movimiento original (ej. "24/08 ROJAS")
        for m in mayor:
            f = FECHA_EN_CONCEPTO.match(m['concepto'])
            if m['ok'] or not f or int(f.group(2)) != banco_ant[0]['fecha'].month:
                continue
            cands = [b for b in banco_ant if not b['ok'] and b['importe'] == m['importe']]
            cands.sort(key=lambda b: abs(b['fecha'].day - int(f.group(1))))
            if cands:
                cerrar('Mes anterior', [m], cands[:1], 'El banco lo registró el mes anterior')
        cola = [b for b in banco_ant if (inicio - b['fecha']).days <= DIAS_CRUCE_MES]
        for m in mayor:
            if m['ok'] or (m['fecha'] - inicio).days >= DIAS_CRUCE_MES:
                continue
            for b in sorted(cola, key=lambda b: (inicio - b['fecha']).days):
                if not b['ok'] and b['importe'] == m['importe']:
                    cerrar('Mes anterior', [m], [b], 'El banco lo registró el mes anterior')
                    break

    # 4) Movimientos que se anulan entre sí dentro del mismo listado (reversas)
    for lista, tipo in ((banco, 'Anulados en banco'), (mayor, 'Anulados en mayor')):
        for a in lista:
            if a['ok']:
                continue
            for c in lista:
                f = FECHA_EN_CONCEPTO.match(c['concepto']) if lista is mayor else None
                referido = bool(f) and (int(f.group(1)), int(f.group(2))) == (a['fecha'].day, a['fecha'].month)
                if (not c['ok'] and c is not a and c['importe'] == -a['importe'] and c['importe'] != 0
                        and (abs((c['fecha'] - a['fecha']).days) <= 3 or referido)):
                    if lista is banco:
                        cerrar(tipo, [], [a, c], 'Movimiento y su reversa')
                    else:
                        cerrar(tipo, [a, c], [], 'Asiento y su contraasiento')
                    break

    # 5) Asiento global de gastos bancarios contra la suma de gastos del resumen
    for m in mayor:
        if m['ok'] or m['importe'] >= 0 or not ASIENTO_GASTOS.search(m['concepto']):
            continue
        for lista, tipo in ((banco, 'Gastos bancarios'), (banco_ant, 'Gastos mes anterior')):
            gs = [b for b in lista if not b['ok'] and rubro_gasto(b['concepto'])]
            dif = round(m['importe'] - sum(b['importe'] for b in gs), 2)
            if gs and abs(dif) <= abs(m['importe']) * TOL_GASTOS:
                nota = f'Suma de {len(gs)} gastos del resumen'
                if dif:
                    nota += f'; el asiento difiere en {dif:,.2f} de esa suma'
                cerrar(tipo, [m], gs, nota)
                break
    return grupos


# ---------------------------------------------------------------- Excel

NUM = '#,##0.00;[Red]-#,##0.00'
FEC = 'dd/mm/yyyy'
TITULO = Font(bold=True, size=13)
NEGRITA = Font(bold=True)
CABECERA = PatternFill('solid', fgColor='D9E1F2')
VERDE = PatternFill('solid', fgColor='E2EFDA')
ROJO = PatternFill('solid', fgColor='FCE4D6')


def hoja(wb, nombre, columnas, filas, anchos):
    ws = wb.create_sheet(nombre)
    ws.append(columnas)
    for c in ws[1]:
        c.font, c.fill = NEGRITA, CABECERA
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for f in filas:
        ws.append(f)
    for i, a in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = a
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            if isinstance(c.value, date):
                c.number_format = FEC
            elif isinstance(c.value, float):
                c.number_format = NUM
    ws.freeze_panes = 'A2'
    return ws


def fila_total(ws, etiqueta, col, total):
    ws.append([])
    f = ws.max_row + 1
    ws.cell(f, 1, etiqueta).font = NEGRITA
    c = ws.cell(f, col, total)
    c.font, c.number_format = NEGRITA, NUM


def cuadre(im, ib, mayor, banco, grupos):
    """Pendientes de cada lado y control de saldos al inicio y al cierre."""
    pend_m = [m for m in mayor if not m['ok']]
    pend_b = [b for b in banco if not b['ok'] and b['importe'] != 0]
    pend_gastos = [b for b in pend_b if rubro_gasto(b['concepto'])]
    pend_otros = [b for b in pend_b if not rubro_gasto(b['concepto'])]
    mes_ant = [g for g in grupos if g['tipo'] in ('Mes anterior', 'Gastos mes anterior')]
    tot_mes_ant = round(sum(m['importe'] for g in mes_ant for m in g['mayor']), 2)
    tot_pend_b = round(sum(b['importe'] for b in pend_b), 2)
    tot_pend_m = round(sum(m['importe'] for m in pend_m), 2)
    tot_gastos = round(sum(b['importe'] for b in pend_gastos), 2)
    difs = round(sum(sum(m['importe'] for m in g['mayor']) - sum(b['importe'] for b in g['banco'])
                     for g in grupos if g['mayor'] and g['banco'] and g not in mes_ant), 2)

    dif_ini = round(ib['saldo_inicial'] - im['saldo_inicial'], 2)
    sin_explicar_ini = round(dif_ini - tot_mes_ant, 2)
    calculado = round(im['saldo_final'] + tot_pend_b - tot_pend_m - difs, 2)
    dif_final = round(ib['saldo_final'] - calculado, 2)
    return locals()


def valores_a_acreditar(banco):
    """Depósitos de valores que el banco todavía no acreditó (informativo)."""
    acreditados = [b for b in banco if b['importe'] > 0 and VALORES.search(b['concepto'])]
    usados, pendientes = set(), []
    for lote in [b for b in banco if b['por_acreditar']]:
        cands = [(round(b['importe'] * 100), b) for b in acreditados
                 if id(b) not in usados and 0 <= (b['fecha'] - lote['fecha']).days <= 6]
        obj = round(lote['por_acreditar'] * 100)
        sel = [c[1] for c in cands if c[0] == obj][:1] or subconjunto(obj, cands)
        if sel:
            usados.update(id(b) for b in sel)
        else:
            pendientes.append(lote)
    return pendientes


def escribir_excel(path, banco_nombre, periodo, im, ib, mayor, banco, banco_ant, grupos, hay_anterior):
    wb = Workbook()
    wb.remove(wb.active)
    c = cuadre(im, ib, mayor, banco, grupos)
    pend_m, pend_otros, pend_gastos, mes_ant = c['pend_m'], c['pend_otros'], c['pend_gastos'], c['mes_ant']
    tot_mes_ant, tot_pend_b, tot_pend_m, tot_gastos = c['tot_mes_ant'], c['tot_pend_b'], c['tot_pend_m'], c['tot_gastos']
    difs, dif_ini, sin_explicar_ini = c['difs'], c['dif_ini'], c['sin_explicar_ini']
    calculado, dif_final = c['calculado'], c['dif_final']

    # --- Resumen
    ws = wb.create_sheet('Resumen')
    ws.column_dimensions['A'].width = 62
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 46

    def linea(texto, valor=None, nota='', negrita=False, relleno=None):
        ws.append([texto, valor, nota])
        f = ws.max_row
        if negrita:
            ws.cell(f, 1).font = ws.cell(f, 2).font = NEGRITA
        if isinstance(valor, float):
            ws.cell(f, 2).number_format = NUM
        if relleno:
            for c in ws[f]:
                c.fill = relleno

    ws.append([f'Conciliación bancaria - {banco_nombre} - {periodo}'])
    ws['A1'].font = TITULO
    linea(f"Cuenta del mayor: {im['cuenta']}")
    ws.append([])
    linea('MOVIMIENTOS', negrita=True)
    linea('Asientos en el mayor', len(mayor))
    linea('Movimientos en el resumen del banco', len([b for b in banco if b['importe'] != 0]))
    linea('Asientos del mayor conciliados', len(mayor) - len(pend_m))
    linea('Asientos del mayor sin conciliar', len(pend_m), 'Ver hoja "En mayor, no en banco"')
    linea('Movimientos del banco sin contabilizar (sin gastos)', len(pend_otros), 'Ver hoja "En banco, no en mayor"')
    linea('Gastos e impuestos del banco sin contabilizar', len(pend_gastos), 'Ver hoja "Gastos del banco"')
    ws.append([])
    linea('SALDOS AL INICIO', negrita=True)
    linea('Saldo inicial según banco', ib['saldo_inicial'])
    linea('Saldo inicial según mayor', im['saldo_inicial'])
    linea('Diferencia al inicio', dif_ini)
    linea('  Explicada por asientos de este mes que el banco registró el mes anterior', tot_mes_ant,
          'Ver hoja "Mes anterior"' if hay_anterior else 'No se encontró el resumen del mes anterior')
    linea('  Diferencia inicial sin explicar', sin_explicar_ini, negrita=True,
          relleno=VERDE if abs(sin_explicar_ini) < 0.005 else ROJO)
    ws.append([])
    linea('SALDOS AL CIERRE', negrita=True)
    linea('Saldo final según mayor', im['saldo_final'])
    linea('  (+/-) Movimientos del banco sin contabilizar', round(tot_pend_b - tot_gastos, 2))
    linea('  (+/-) Gastos e impuestos del banco sin contabilizar', tot_gastos)
    linea('  (-/+) Asientos del mayor que no están en el banco', round(-tot_pend_m, 2))
    linea('  (-/+) Diferencias de importe en partidas conciliadas', round(-difs, 2))
    linea('Saldo del banco calculado', calculado, negrita=True)
    linea('Saldo final según banco', ib['saldo_final'], negrita=True)
    linea('Diferencia final', dif_final,
          'Debe coincidir con la diferencia inicial sin explicar', negrita=True,
          relleno=VERDE if abs(dif_final) < 0.005 else ROJO)
    con_dif = [g for g in grupos if 'ifiere' in g['nota'] or 'Diferencia' in g['nota']]
    if con_dif:
        ws.append([])
        linea('PARTIDAS CONCILIADAS CON DIFERENCIA DE IMPORTE', negrita=True)
        for g in con_dif:
            m = g['mayor'][0]
            linea(f"  {m['fecha']:%d/%m/%Y} {m['concepto']} (asiento {m['nro']})", m['importe'], g['nota'])

    # --- Pendientes
    ws = hoja(wb, 'En mayor, no en banco',
              ['Fecha', 'Concepto', 'Nro', 'Suc', 'Debe', 'Haber'],
              [[m['fecha'], m['concepto'], m['nro'], m['suc'],
                m['importe'] if m['importe'] > 0 else None, -m['importe'] if m['importe'] < 0 else None]
               for m in pend_m], [12, 34, 8, 6, 16, 16])
    fila_total(ws, 'Neto (debe - haber)', 5, tot_pend_m)

    ws = hoja(wb, 'En banco, no en mayor',
              ['Fecha', 'Concepto', 'Comprob.', 'Crédito', 'Débito', 'Detalle'],
              [[b['fecha'], b['concepto'], b['comprob'],
                b['importe'] if b['importe'] > 0 else None, -b['importe'] if b['importe'] < 0 else None,
                b['detalle']] for b in pend_otros], [12, 44, 13, 16, 16, 70])
    fila_total(ws, 'Neto (crédito - débito)', 4, round(tot_pend_b - tot_gastos, 2))

    rubros = {}
    for b in pend_gastos:
        r = rubros.setdefault(rubro_gasto(b['concepto']), [0, 0.0])
        r[0] += 1
        r[1] = round(r[1] - b['importe'], 2)
    ws = hoja(wb, 'Gastos del banco', ['Rubro', 'Cantidad', 'Importe', '', 'Fecha', 'Concepto', 'Importe'],
              [], [34, 10, 16, 3, 12, 44, 16])
    detalle = [[b['fecha'], b['concepto'], -b['importe']] for b in pend_gastos]
    resumen = [[k, v[0], v[1]] for k, v in rubros.items()] + [['TOTAL', len(pend_gastos), -tot_gastos]]
    for r, d in zip_longest(resumen, detalle, fillvalue=[None, None, None]):
        ws.append(r + [None] + d)
    for fila in ws.iter_rows(min_row=2):
        fila[2].number_format = fila[6].number_format = NUM
        fila[4].number_format = FEC
        if fila[0].value == 'TOTAL':
            for c in fila[:3]:
                c.font = NEGRITA

    # --- Mes anterior
    filas = []
    for g in mes_ant:
        m = g['mayor'][0]
        b0 = g['banco'][0]
        filas.append([m['fecha'], m['concepto'], m['nro'], m['importe'],
                      b0['fecha'] if len(g['banco']) == 1 else g['banco'][-1]['fecha'],
                      b0['concepto'] if len(g['banco']) == 1 else g['nota'],
                      round(sum(b['importe'] for b in g['banco']), 2),
                      round(m['importe'] - sum(b['importe'] for b in g['banco']), 2) or None])
    ws = hoja(wb, 'Mes anterior',
              ['Fecha mayor', 'Concepto mayor', 'Nro', 'Importe mayor',
               'Fecha banco', 'Concepto banco', 'Importe banco', 'Diferencia'],
              filas, [12, 34, 8, 16, 12, 70, 16, 14])
    fila_total(ws, 'Total', 4, tot_mes_ant)

    # --- Conciliados
    filas = []
    n = 0
    for g in grupos:
        if g in mes_ant:
            continue
        n += 1
        bs = g['banco']
        if g['tipo'] == 'Gastos bancarios':
            bs = [{'fecha': bs[-1]['fecha'], 'concepto': g['nota'],
                   'importe': round(sum(b['importe'] for b in bs), 2)}]
        for i, (m, b) in enumerate(zip_longest(g['mayor'], bs)):
            filas.append([n, g['tipo'] if i == 0 else None,
                          m and m['fecha'], m and m['concepto'], m and m['nro'], m and m['importe'],
                          b and b['fecha'], b and b['concepto'], b and b['importe'],
                          g['nota'] if i == 0 and g['tipo'] != 'Gastos bancarios' else None])
    hoja(wb, 'Conciliados',
         ['#', 'Tipo', 'Fecha mayor', 'Concepto mayor', 'Nro', 'Importe mayor',
          'Fecha banco', 'Concepto banco', 'Importe banco', 'Observación'],
         filas, [5, 20, 12, 30, 7, 16, 12, 44, 16, 44])

    filas = [[l['fecha'], l['concepto'], l['comprob'], l['por_acreditar'], l['detalle']]
             for l in valores_a_acreditar(banco)]
    hoja(wb, 'Valores a acreditar', ['Fecha depósito', 'Concepto', 'Comprob.', 'Importe', 'Detalle'],
         filas, [14, 26, 13, 16, 70])

    wb.save(path)
    return {'pend_mayor': len(pend_m), 'pend_banco': len(pend_otros), 'gastos': len(pend_gastos),
            'conciliados': len(mayor) - len(pend_m), 'sin_explicar_ini': sin_explicar_ini,
            'dif_final': dif_final}


def archivos(nombre, mes, anio):
    """Todos los PDF que intervienen: mayor, resúmenes del mes y del mes anterior."""
    mes_ant, anio_ant = (mes - 1, anio) if mes > 1 else (12, anio - 1)
    return ([BASE / 'MAYOR' / f'{nombre}{mes:02d}{anio}.pdf']
            + resumenes(nombre, BASE, mes, anio) + resumenes(nombre, BASE, mes_ant, anio_ant))


def cargar(nombre, mes, anio):
    """Lee los PDF y concilia. Devuelve (im, ib, mayor, banco, banco_ant, grupos, hay_anterior)."""
    f_mayor = BASE / 'MAYOR' / f'{nombre}{mes:02d}{anio}.pdf'
    if not f_mayor.exists():
        raise ErrorLectura(f'No encuentro el mayor: {f_mayor}')
    im, mayor = leer_mayor(f_mayor)
    actual = leer_resumen(nombre, BASE, mes, anio)
    if actual is None:
        raise ErrorLectura(f"No encuentro el resumen de {nombre} de {mes:02d}/{anio} en la carpeta "
                           f"{BANCOS[nombre]['carpeta']}.")
    ib, banco = actual
    mes_ant, anio_ant = (mes - 1, anio) if mes > 1 else (12, anio - 1)
    anterior = leer_resumen(nombre, BASE, mes_ant, anio_ant)
    banco_ant = anterior[1] if anterior else []
    grupos = conciliar(mayor, banco, banco_ant)
    return im, ib, mayor, banco, banco_ant, grupos, anterior is not None


def main(argv):
    if len(argv) != 3 or argv[1].upper() not in BANCOS or not re.fullmatch(r'\d{1,2}/\d{4}', argv[2]):
        print(__doc__)
        print('Bancos disponibles:', ', '.join(BANCOS))
        return 1
    nombre = argv[1].upper()
    mes, anio = (int(x) for x in argv[2].split('/'))
    try:
        im, ib, mayor, banco, banco_ant, grupos, hay_ant = cargar(nombre, mes, anio)
    except ErrorLectura as e:
        print('ERROR:', e)
        return 1

    salida = BASE / 'CONCILIACIONES'
    salida.mkdir(exist_ok=True)
    destino = salida / f'{nombre} {mes:02d}-{anio}.xlsx'
    r = escribir_excel(destino, nombre, f'{mes:02d}/{anio}', im, ib, mayor, banco, banco_ant, grupos, hay_ant)

    print(f'{nombre} {mes:02d}/{anio}')
    print(f"  Asientos del mayor conciliados:          {r['conciliados']} de {len(mayor)}")
    print(f"  En mayor y no en banco:                  {r['pend_mayor']}")
    print(f"  En banco y no en mayor (sin gastos):     {r['pend_banco']}")
    print(f"  Gastos/impuestos del banco sin asiento:  {r['gastos']}")
    print(f"  Diferencia inicial sin explicar:         {r['sin_explicar_ini']:,.2f}")
    print(f"  Diferencia final:                        {r['dif_final']:,.2f}")
    print(f'  Resultado: {destino}')
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv))
